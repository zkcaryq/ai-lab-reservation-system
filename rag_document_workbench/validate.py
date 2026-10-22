# -*- coding: utf-8 -*-
"""最终验证：打开性、图片嵌入、文本层、事实一致性、违规内容、错误值。"""
import io
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import load_workbook
from docx import Document

ROOT = "D:/AI/faceprogram/rag_source_documents"
report = {"docx": 0, "pdf": 0, "xlsx": 0, "md": 0, "issues": []}


def walk():
    for dirpath, dirnames, filenames in os.walk(ROOT):
        for fn in filenames:
            yield os.path.join(dirpath, fn)


def check_docx(path):
    try:
        doc = Document(path)
        n_imgs = 0
        with zipfile.ZipFile(path) as z:
            n_imgs = len([n for n in z.namelist() if n.startswith("word/media/")])
        report["docx"] += 1
        return n_imgs
    except Exception as e:
        report["issues"].append(f"DOCX打不开: {path}: {e}")
        return -1


def check_pdf(path):
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(path)
        n = len(pdf)
        chars = 0
        for i in range(min(n, 3)):
            chars += len(pdf[i].get_textpage().get_text_range().strip())
        report["pdf"] += 1
        return n, chars
    except Exception as e:
        report["issues"].append(f"PDF打不开: {path}: {e}")
        return -1, -1


def check_xlsx(path):
    try:
        wb = load_workbook(path)
        errors = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("#") and cell.value.endswith(("!", "?")):
                        errors.append(f"{os.path.basename(path)}:{ws.title}:{cell.coordinate}:{cell.value}")
        report["xlsx"] += 1
        return len(wb.worksheets), errors
    except Exception as e:
        report["issues"].append(f"XLSX打不开: {path}: {e}")
        return -1, [str(e)]


FACT_CHECKS = [
    ("39 人", "03-机器人创新实验室-综合介绍与使用手册.docx", "容量39"),
    ("设备维护", "13-综合开放实验室-综合介绍与使用手册.docx", "综合开放关闭原因"),
    ("维修中", "40-3D打印机-操作手册.docx", "3D打印机维修中"),
]


def main():
    docx_imgs = {}
    pdf_info = {}
    xlsx_errs = []
    secret_pat = re.compile(r"zkcaryq770880|sk-ws-H\.PYYLPRR|qc125F0u2bZB6", re.I)
    secret_hits = []
    phone_pat = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")

    for p in walk():
        ext = p.rsplit(".", 1)[-1].lower()
        rel = os.path.relpath(p, ROOT)
        if ext == "docx":
            n = check_docx(p)
            docx_imgs[rel] = n
        elif ext == "pdf":
            n, chars = check_pdf(p)
            pdf_info[rel] = (n, chars)
        elif ext == "xlsx":
            _, errs = check_xlsx(p)
            xlsx_errs += errs
        elif ext == "md":
            report["md"] += 1
        # 内容扫描
        if ext in ("docx", "xlsx", "md", "pdf"):
            raw = b""
            if ext in ("docx", "xlsx"):
                with zipfile.ZipFile(p) as z:
                    for n in z.namelist():
                        if n.endswith(".xml"):
                            raw += z.read(n)
            elif ext == "md":
                raw = open(p, "rb").read()
            if secret_pat.search(raw.decode("utf-8", "ignore")):
                secret_hits.append(rel)
            if ext != "07" and rel.startswith("0"):
                for m in phone_pat.finditer(raw.decode("utf-8", "ignore")):
                    num = m.group()
                    if not num.startswith("13800001") and not num.startswith("1380000100") and not num.startswith("0731"):
                        report["issues"].append(f"疑似真实手机号: {rel}: {num}")

    # 事实一致性（文本抽查）
    for pat, fname, label in FACT_CHECKS:
        p = None
        for dirpath, _, fns in os.walk(ROOT):
            for fn in fns:
                if fn == fname:
                    p = os.path.join(dirpath, fn)
        if not p:
            report["issues"].append(f"事实检查找不到文件: {fname}")
            continue
        doc = Document(p)
        text = "\n".join(par.text for par in doc.paragraphs)
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    text += "\n" + c.text
        if not re.search(pat, text):
            report["issues"].append(f"事实检查失败: {label}（{fname}）")
        else:
            print("FACT OK:", label)

    # AI实验室脏数字检查
    p1 = f"{ROOT}/01-实验室介绍与使用手册/01-人工智能实验室-综合介绍与使用手册.docx"
    doc = Document(p1)
    text = "\n".join(par.text for par in doc.paragraphs)
    if re.search(r"543543|6666666666", text):
        report["issues"].append("AI实验室手册包含脏数字！")
    else:
        print("FACT OK: AI实验室简介无脏数字")

    # Markdown 相对路径
    mddir = f"{ROOT}/08-Markdown通知与指南"
    for fn in os.listdir(mddir):
        if fn.endswith(".md"):
            for m in re.finditer(r"\]\((\.media/[^)]+)\)", open(f"{mddir}/{fn}", encoding="utf-8").read()):
                target = f"{mddir}/{m.group(1)}"
                if not os.path.exists(target):
                    report["issues"].append(f"MD图片路径失效: {fn}: {m.group(1)}")
    print("MD图片路径检查完成")

    # 禁止出现的加工目录
    for bad in ["chunks", "embeddings", "parsed", "extracted_images", "ocr_results",
                "vector_data", "assets", "tools"]:
        if os.path.exists(f"{ROOT}/{bad}"):
            report["issues"].append(f"发现禁止目录: {bad}")
    print("禁止目录检查完成")

    # 统计含图docx
    with_img = [k for k, v in docx_imgs.items() if v > 0]
    print(f"\nDOCX总数: {report['docx']}（含图 {len(with_img)}）")
    print(f"PDF总数: {report['pdf']}")
    print(f"XLSX总数: {report['xlsx']}，公式错误: {len(xlsx_errs)}")
    print(f"MD总数: {report['md']}")
    print(f"密钥泄露: {len(secret_hits)}")
    for e in xlsx_errs[:5]:
        print("XLSX错误示例:", e)
    print(f"\n问题清单（{len(report['issues'])}）:")
    for i in report["issues"][:20]:
        print(" -", i)
    if not report["issues"]:
        print("（无问题）")


if __name__ == "__main__":
    main()
