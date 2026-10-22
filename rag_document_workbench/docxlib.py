# -*- coding: utf-8 -*-
"""python-docx 辅助库：统一版式（封面、标题、表格、图片、页眉页脚）。"""
import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

IMG = "D:/AI/faceprogram/rag_document_workbench/gen_images"
PHOTOS = "D:/AI/faceprogram/rag_document_workbench/photos"
DARK = RGBColor(0x1A, 0x3A, 0x6B)
GREY = RGBColor(0x5F, 0x63, 0x68)

_fonts = {"hei": "黑体", "song": "宋体", "kai": "楷体", "yahei": "微软雅黑"}


def set_run(run, size=12, bold=False, color=None, font="yahei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = _fonts.get(font, font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), _fonts.get(font, font))


def new_doc(landscape=False):
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        from docx.enum.section import WD_ORIENT
        w, h = sec.page_width, sec.page_height
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = h, w
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(11)
    return doc


def cover(doc, title, subtitle, org, meta_lines, img_path=None):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(title), size=30, bold=True, color=DARK)
    if subtitle:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(subtitle), size=16, color=GREY)
    if img_path and os.path.exists(img_path):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img_path, width=Cm(13.5))
    for _ in range(2):
        doc.add_paragraph()
    for line in meta_lines:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(line), size=12, color=GREY)
    doc.add_page_break()


def h1(doc, text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    set_run(p.add_run((f"{num} " if num else "") + text), size=16, bold=True, color=DARK)
    return p


def h2(doc, text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run((f"{num} " if num else "") + text), size=13, bold=True)
    return p


def para(doc, text, size=11, bold=False, color=None, align=None, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def bullets(doc, items, size=11):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        set_run(p.add_run("• " + it), size=size)


def numbered(doc, items, size=11, start=1):
    for i, it in enumerate(items, start):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        set_run(p.add_run(f"{i}. {it}"), size=size)


def kv_table(doc, rows, widths=None, header=None):
    """rows: list of tuple(str)。header: list[str] 或 None"""
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows) + (1 if header else 0), cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    ri = 0
    if header:
        for ci, htxt in enumerate(header):
            cell = t.cell(0, ci)
            cell.paragraphs[0].text = ""
            set_run(cell.paragraphs[0].add_run(htxt), size=10.5, bold=True, color=DARK)
        ri = 1
    for row in rows:
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.paragraphs[0].text = ""
            set_run(cell.paragraphs[0].add_run(str(val)), size=10.5)
        ri += 1
    if widths:
        for ci, w in enumerate(widths):
            for r in t.rows:
                r.cells[ci].width = Cm(w)
    return t


def pic(doc, path, width_cm=13.0, caption=None, center=True):
    if not os.path.exists(path):
        para(doc, f"[图片缺失：{os.path.basename(path)}]", color=RGBColor(0xCC, 0, 0))
        return
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        cp = doc.add_paragraph()
        if center:
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(cp.add_run(caption), size=9.5, color=GREY)


def footer_note(doc, text):
    sec = doc.sections[0]
    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(f.add_run(text), size=8.5, color=GREY)


def revision_table(doc, version, date, rows):
    h1(doc, "文档版本与修订记录")
    all_rows = [(version, date, r[0], r[1]) for r in rows]
    kv_table(doc, all_rows, widths=[2.2, 2.8, 7.0, 3.2], header=["版本", "日期", "修订内容", "编制/审核"])
