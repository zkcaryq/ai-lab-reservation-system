"""把DOCX、PDF、XLSX、Markdown和图片统一解析成ParsedChunk。"""

import asyncio
import logging
import re
import shutil
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from app.core.errors import BusinessException
from app.core.settings import get_settings
from app.services.knowledge.bailian_document_services import analyze_image

settings = get_settings()
logger = logging.getLogger(__name__)

# Markdown图片语法中的圆括号部分，例如``![平面图](images/lab.png)``。
# 第一版只处理本地相对路径；HTTP远程图片不会自动下载，避免服务端请求伪造风险。
MARKDOWN_IMAGE_PATTERN = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")

# 只有这些常见位图格式能够同时交给OCR、视觉理解和融合向量接口。
# DOCX/PDF内部可能包含EMF、WMF、SVG等资源；第一版跳过它们，避免整份文档入库失败。
MODEL_IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}

# 中文制度、手册以字符数控制比引入特定模型 tokenizer 更稳定。800 字是“软目标”：
# 达到目标后只在自然语言边界换块，不能为了凑固定长度把一句话从中间切断。
TEXT_CHUNK_TARGET_CHARS = 800
SECTION_HEADING_PATTERN = re.compile(
    r"^(?:第[一二三四五六七八九十百\d]+[章节条部分]|"
    r"[一二三四五六七八九十]+、)"
)
SENTENCE_BOUNDARY_PATTERN = re.compile(r"(?<=[。！？；!?;])")
CLAUSE_BOUNDARY_PATTERN = re.compile(r"(?<=[，、,:：])")


@dataclass(slots=True)
class ParsedChunk:
    """不同文件解析器统一交给入库Service的中间数据。

    这个类不是数据库Model。DOCX、PDF、Excel、Markdown和图片虽然来源不同，
    最终都会先转换为ParsedChunk，后续向量化代码因此不必为每种格式各写一套。
    """

    # content既用于生成向量，也会在RAG检索命中后交给聊天模型。
    content: str

    # text表示纯文字块；image表示文字与图片一起生成融合向量。
    modality: str = "text"

    # image_path只在处理期间使用，指向后端私有目录，不会直接返回前端。
    image_path: Path | None = None

    # image_storage_key是可持久化的相对路径，数据库不能保存本机绝对路径。
    image_storage_key: str | None = None
    page_number: int | None = None
    sheet_name: str | None = None
    heading_path: str | None = None
    source_metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class _PdfImageAsset:
    """同步PDF提取阶段交给异步模型阶段的内部图片记录。"""

    path: Path
    page_number: int
    surrounding_text: str
    is_scan_page: bool
    image_index: int | None = None


def _write_working_image(raw: bytes, suffix: str, *, working_dir: Path) -> Path:
    """把提取图片写入本次任务独享的临时目录。"""

    # UUID服务器文件名避免原始文件名冲突，也不信任文档内部携带的路径。
    path = working_dir / f"{uuid.uuid4().hex}{suffix.lower()}"

    # 此目录由FastAPI lifespan在Worker启动前创建。
    path.write_bytes(raw)
    return path


def _image_content(
    *,
    ocr_text: str | None,
    description: str | None,
    context_label: str | None = None,
    context: str = "",
) -> str:
    """只组合实际成功的图片分析结果，不写入虚假的占位内容。"""

    sections: list[str] = []
    if ocr_text:
        sections.append(f"图片OCR：\n{ocr_text}")
    if description:
        sections.append(f"图片说明：\n{description}")
    if context_label and context:
        sections.append(f"{context_label}：\n{context}")
    return "\n\n".join(sections)


def _split_text(text: str, *, heading: str | None = None) -> list[ParsedChunk]:
    """按标题、段落和句子分层切块，目标约800字符。

    旧实现只按空行切分；DOCX解析后段落之间通常只有一个换行，而且单个大段
    超过上限时也不会继续拆，实际出现过3000多字的Chunk。这里先保留自然段，
    再按完整句子拆分。单句超过软目标时，只允许在逗号、冒号等从句边界继续拆；
    如果原文自身没有任何标点，就保留这个长句，不能在固定字符位置制造断句。
    """

    # 单换行也代表DOCX段落或PDF文本行；连续空白行不会产生空块。
    raw_blocks = [line.strip() for line in re.split(r"\n+", text) if line.strip()]
    units: list[str] = []
    for block in raw_blocks:
        if len(block) <= TEXT_CHUNK_TARGET_CHARS:
            units.append(block)
            continue

        # 超长段优先按完整句子拆，句末标点会保留在前一个单元中。
        for sentence in SENTENCE_BOUNDARY_PATTERN.split(block):
            sentence = sentence.strip()
            if not sentence:
                continue
            if len(sentence) <= TEXT_CHUNK_TARGET_CHARS:
                units.append(sentence)
                continue

            # 一个语法句可能包含多个很长的并列分句。这里按逗号、顿号、冒号等
            # 自然边界重新组合，而不是使用 sentence[:800] 从词语中间切断。
            clauses = [value.strip() for value in CLAUSE_BOUNDARY_PATTERN.split(sentence)]
            clause_buffer = ""
            for clause in clauses:
                if not clause:
                    continue
                if clause_buffer and len(clause_buffer) + len(clause) > TEXT_CHUNK_TARGET_CHARS:
                    units.append(clause_buffer)
                    clause_buffer = ""
                clause_buffer += clause
            if clause_buffer:
                units.append(clause_buffer)

    chunks: list[ParsedChunk] = []
    buffer: list[str] = []
    buffer_length = 0

    def flush() -> None:
        """把当前缓冲区固化为一个完整文本 Chunk。"""

        nonlocal buffer_length
        if not buffer:
            return
        chunks.append(ParsedChunk(content="\n".join(buffer), heading_path=heading))
        buffer.clear()
        buffer_length = 0

    for unit in units:
        # 章节标题不能黏在上一章末尾，否则检索命中后上下文语义会混杂。
        if buffer and SECTION_HEADING_PATTERN.match(unit):
            flush()
        separator_length = 1 if buffer else 0
        if buffer and buffer_length + separator_length + len(unit) > TEXT_CHUNK_TARGET_CHARS:
            flush()
            separator_length = 0
        buffer.append(unit)
        buffer_length += separator_length + len(unit)
    flush()
    return chunks


def _parse_docx_sync(path: Path, working_dir: Path) -> tuple[list[ParsedChunk], list[Path]]:
    """在线程中读取DOCX文字、表格和可处理的内嵌图片。"""

    # python-docx负责读取段落与表格；DOCX本质是ZIP，内嵌图片稍后从word/media提取。
    document = Document(path)
    text_parts: list[str] = []
    heading = ""
    chunks: list[ParsedChunk] = []

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            value = item.text.strip()
            if not value:
                continue
            if item.style and item.style.name.startswith("Heading"):
                if text_parts:
                    chunks.extend(_split_text("\n".join(text_parts), heading=heading or None))
                    text_parts.clear()
                heading = value
            else:
                text_parts.append(value)
        elif isinstance(item, Table):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in item.rows]
            text_parts.append("\n".join(rows))

    if text_parts:
        chunks.extend(_split_text("\n".join(text_parts), heading=heading or None))

    images: list[Path] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/") or name.endswith("/"):
                continue
            suffix = Path(name).suffix.lower()
            if suffix not in MODEL_IMAGE_SUFFIXES:
                # 例如Office中的EMF/WMF矢量图当前模型链路不能稳定读取，先跳过。
                continue
            if len(images) >= settings.KNOWLEDGE_MAX_IMAGES:
                raise ValueError("文档中的图片数量超过处理上限")
            image_path = _write_working_image(
                archive.read(name),
                suffix,
                working_dir=working_dir,
            )
            images.append(image_path)
    return chunks, images


async def parse_docx(path: Path, *, working_dir: Path) -> list[ParsedChunk]:
    """解析DOCX，并对有意义的内嵌图片并发执行OCR与视觉理解。"""

    # ZIP解包和python-docx解析都是同步磁盘/CPU操作，放到线程避免阻塞事件循环。
    chunks, images = await asyncio.to_thread(_parse_docx_sync, path, working_dir)
    document_context = chunks[0].content[:1000] if chunks else path.stem
    for image_path in images:
        # 小于10KB通常是图标、分隔线或低价值装饰，第一版跳过模型调用。
        if image_path.stat().st_size < 10 * 1024:
            # 文件已经由同步解析阶段复制到私有目录；既然不入库，就立即清理。
            image_path.unlink(missing_ok=True)
            continue
        try:
            ocr_text, description = await analyze_image(
                image_path,
                surrounding_text=document_context,
            )
        except BusinessException:
            # DOCX正文已经可以入库时，单张装饰图或异常图不应拖垮整份文档。
            logger.warning("skipping unreadable DOCX image image=%s", image_path.name)
            image_path.unlink(missing_ok=True)
            continue
        chunks.append(
            ParsedChunk(
                content=_image_content(
                    ocr_text=ocr_text,
                    description=description,
                    context_label="文档上下文",
                    context=document_context,
                ),
                modality="image",
                image_path=image_path,
            )
        )
    return chunks


def _extract_pdf_sync(
    path: Path,
    working_dir: Path,
) -> tuple[list[ParsedChunk], list[_PdfImageAsset]]:
    """在线程中完成全部PyMuPDF页面读取、渲染和图片提取。"""

    document = pymupdf.open(path)
    chunks: list[ParsedChunk] = []
    assets: list[_PdfImageAsset] = []
    try:
        if document.page_count > settings.KNOWLEDGE_MAX_PAGES:
            raise ValueError("PDF页数超过处理上限")
        image_count = 0
        for page_index in range(document.page_count):
            page = document.load_page(page_index)
            text = page.get_text("text").strip()

            # 有足够文本层时直接使用，避免重复OCR。
            if len(text) >= 20:
                for chunk in _split_text(text):
                    chunk.page_number = page_index + 1
                    chunks.append(chunk)
            else:
                # 扫描页渲染为PNG，再做OCR和视觉理解。
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                image_count += 1
                if image_count > settings.KNOWLEDGE_MAX_IMAGES:
                    raise ValueError("PDF中的图片数量超过处理上限")
                image_path = _write_working_image(
                    pixmap.tobytes("png"),
                    ".png",
                    working_dir=working_dir,
                )
                assets.append(
                    _PdfImageAsset(
                        path=image_path,
                        page_number=page_index + 1,
                        surrounding_text="",
                        is_scan_page=True,
                    )
                )

                # 扫描页本身通常就是一张整页图片，已经完成OCR和视觉理解。
                # 继续page.get_images()会重复处理同一扫描图，造成重复Chunk和双倍模型费用。
                continue

            # 原生PDF中的业务图片单独提取，用于平面图和设备图检索。
            for image_index, image_info in enumerate(page.get_images(full=True)):
                image = document.extract_image(image_info[0])
                raw = image["image"]
                if len(raw) < 10 * 1024:
                    continue
                suffix = f".{image['ext']}"
                if suffix.lower() not in MODEL_IMAGE_SUFFIXES:
                    continue
                image_count += 1
                if image_count > settings.KNOWLEDGE_MAX_IMAGES:
                    raise ValueError("PDF中的图片数量超过处理上限")
                image_path = _write_working_image(raw, suffix, working_dir=working_dir)
                assets.append(
                    _PdfImageAsset(
                        path=image_path,
                        page_number=page_index + 1,
                        surrounding_text=text[:1500],
                        is_scan_page=False,
                        image_index=image_index,
                    )
                )
    finally:
        document.close()
    return chunks, assets


async def parse_pdf(path: Path, *, working_dir: Path) -> list[ParsedChunk]:
    """逐页解析PDF：同步PyMuPDF工作在线程中，模型调用保留异步。"""

    # PyMuPDF的load_page/get_text/get_pixmap/extract_image均为同步CPU或文件操作。
    # 整个提取阶段一次性放入线程，避免大PDF阻塞FastAPI事件循环。
    chunks, assets = await asyncio.to_thread(_extract_pdf_sync, path, working_dir)
    for asset in assets:
        try:
            ocr_text, description = await analyze_image(
                asset.path,
                surrounding_text=asset.surrounding_text,
            )
        except BusinessException:
            if asset.is_scan_page:
                # 扫描页没有可替代文本层，两项都失败意味着该页无法入库。
                raise
            logger.warning("skipping unreadable PDF image image=%s", asset.path.name)
            asset.path.unlink(missing_ok=True)
            continue

        chunks.append(
            ParsedChunk(
                content=_image_content(
                    ocr_text=ocr_text,
                    description=description,
                    context_label="页面正文" if asset.surrounding_text else None,
                    context=asset.surrounding_text,
                ),
                modality="image",
                image_path=asset.path,
                page_number=asset.page_number,
                source_metadata=(
                    {"image_index": asset.image_index}
                    if asset.image_index is not None
                    else {"scan_page": True}
                ),
            )
        )
    return chunks


def _parse_xlsx_sync(path: Path) -> list[ParsedChunk]:
    """把Excel每一行转换成带列名的文字Chunk。"""

    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks: list[ParsedChunk] = []
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers_row = next(rows, None)
            if headers_row is None:
                continue
            headers = [str(value).strip() if value is not None else "" for value in headers_row]
            # read_only配合迭代器逐行读取，不再用list()把整张表复制到内存。
            for row_number, row in enumerate(rows, start=2):
                fields = []
                for column_index, value in enumerate(row):
                    if value is None:
                        continue
                    header = headers[column_index] if column_index < len(headers) else ""
                    label = header or f"第{column_index + 1}列"
                    fields.append(f"{label}：{value}")
                if fields:
                    chunks.append(
                        ParsedChunk(
                            content="；".join(fields),
                            sheet_name=worksheet.title,
                            source_metadata={"row_number": row_number},
                        )
                    )
    finally:
        workbook.close()
    return chunks


async def parse_xlsx(path: Path) -> list[ParsedChunk]:
    """在线程中运行openpyxl，避免同步读取工作簿阻塞其他HTTP请求。"""

    return await asyncio.to_thread(_parse_xlsx_sync, path)


def _safe_extract_zip(zip_path: Path, target: Path) -> None:
    """解压Markdown文档包，阻止绝对路径、../和压缩炸弹。"""

    max_files = 500
    max_uncompressed = 100 * 1024 * 1024
    with zipfile.ZipFile(zip_path) as archive:
        members = archive.infolist()
        if len(members) > max_files:
            raise ValueError("Markdown文档包文件数量过多")
        if sum(item.file_size for item in members) > max_uncompressed:
            raise ValueError("Markdown文档包解压后体积过大")
        target_root = target.resolve()
        for member in members:
            # ZIP成员名理论上使用/，但恶意文件可能使用Windows反斜杠绕过检查。
            normalized_name = member.filename.replace("\\", "/")
            output = (target_root / normalized_name).resolve()
            if not output.is_relative_to(target_root):
                raise ValueError("Markdown文档包包含非法路径")
            # Unix文件类型位0o120000表示符号链接；私有文档包不接受链接成员。
            if (member.external_attr >> 16) & 0o170000 == 0o120000:
                raise ValueError("Markdown文档包不能包含符号链接")
        archive.extractall(target)


async def _parse_markdown_file(
    markdown_path: Path,
    *,
    root: Path,
    working_dir: Path,
) -> list[ParsedChunk]:
    """解析一份Markdown及其相对图片；root只用于限制图片不能越界。"""

    chunks: list[ParsedChunk] = []
    text = markdown_path.read_text(encoding="utf-8")
    chunks.extend(_split_text(text, heading=markdown_path.stem))
    image_values = MARKDOWN_IMAGE_PATTERN.findall(text)
    if len(image_values) > settings.KNOWLEDGE_MAX_IMAGES:
        raise ValueError("Markdown引用的图片数量超过处理上限")
    for relative_value in image_values:
        relative_path = relative_value.strip().split(maxsplit=1)[0].strip("<>")
        if relative_path.startswith(("http://", "https://")):
            continue
        image_path = (markdown_path.parent / relative_path).resolve()
        if not image_path.is_relative_to(root.resolve()) or not image_path.is_file():
            raise ValueError(f"Markdown引用图片不存在或越界：{relative_path}")
        raw = image_path.read_bytes()
        if image_path.suffix.lower() not in MODEL_IMAGE_SUFFIXES:
            continue
        private_path = _write_working_image(
            raw,
            image_path.suffix,
            working_dir=working_dir,
        )
        try:
            ocr_text, description = await analyze_image(
                private_path,
                surrounding_text=text[:1500],
            )
        except BusinessException:
            logger.warning("skipping unreadable Markdown image image=%s", private_path.name)
            private_path.unlink(missing_ok=True)
            continue
        chunks.append(
            ParsedChunk(
                content=_image_content(ocr_text=ocr_text, description=description),
                modality="image",
                image_path=private_path,
                heading_path=markdown_path.stem,
                source_metadata={"markdown_image": relative_path},
            )
        )
    return chunks


async def _parse_markdown_root(root: Path, *, working_dir: Path) -> list[ParsedChunk]:
    """解析ZIP展开目录中的全部Markdown入口。"""

    chunks: list[ParsedChunk] = []
    for markdown_path in sorted(root.rglob("*.md")):
        chunks.extend(
            await _parse_markdown_file(
                markdown_path,
                root=root,
                working_dir=working_dir,
            )
        )
    return chunks


async def parse_markdown(path: Path, *, working_dir: Path) -> list[ParsedChunk]:
    """解析单独上传的 Markdown；本地关联图片应改用 ZIP 文档包。"""

    # 单独上传.md时只解析这一份文件，不能误扫originals目录中的其他文档。
    # 如果Markdown引用本地图片，应改为上传.zip文档包，使图片随包一起进入私有存储。
    return await _parse_markdown_file(path, root=path.parent, working_dir=working_dir)


async def parse_markdown_zip(path: Path, *, working_dir: Path) -> list[ParsedChunk]:
    """在一次性临时目录中解压并解析Markdown文档包。"""

    # 每次任务使用独立UUID目录，多个Worker不会互相覆盖解压文件。
    working = working_dir / "markdown-package"
    working.mkdir(parents=True)
    try:
        await asyncio.to_thread(_safe_extract_zip, path, working)
        if not any(working.rglob("*.md")):
            raise ValueError("ZIP中没有Markdown入口文件")
        return await _parse_markdown_root(working, working_dir=working_dir)
    finally:
        await asyncio.to_thread(shutil.rmtree, working, True)


async def parse_image(path: Path) -> list[ParsedChunk]:
    """把单独上传的图片转换成一个图文融合Chunk。"""

    ocr_text, description = await analyze_image(path)
    # 原始上传图片已经位于originals。正式版本可复制到images，也可复用该storage_key。
    return [
        ParsedChunk(
            content=_image_content(ocr_text=ocr_text, description=description),
            modality="image",
            image_path=path,
        )
    ]


async def parse_document(
    path: Path,
    *,
    extension: str,
    working_dir: Path,
) -> list[ParsedChunk]:
    """按后缀分发解析器；不支持格式明确失败。"""

    if extension == ".docx":
        chunks = await parse_docx(path, working_dir=working_dir)
    elif extension == ".pdf":
        chunks = await parse_pdf(path, working_dir=working_dir)
    elif extension == ".xlsx":
        chunks = await parse_xlsx(path)
    elif extension == ".md":
        chunks = await parse_markdown(path, working_dir=working_dir)
    elif extension == ".zip":
        chunks = await parse_markdown_zip(path, working_dir=working_dir)
    elif extension in {".jpg", ".jpeg", ".png", ".webp"}:
        chunks = await parse_image(path)
    else:
        raise ValueError(f"没有对应解析器：{extension}")

    # PDF文本层、OCR或Office内部XML偶尔会带NUL控制字符。Python字符串允许它，
    # 但PostgreSQL text/varchar明确拒绝0x00，因此必须在统一解析出口清理。
    for chunk in chunks:
        chunk.content = chunk.content.replace("\x00", "").strip()
    chunks = [chunk for chunk in chunks if chunk.content]

    if len(chunks) > settings.KNOWLEDGE_MAX_CHUNKS:
        raise ValueError("文档分块数量超过处理上限")
    if sum(len(chunk.content) for chunk in chunks) > settings.KNOWLEDGE_MAX_TEXT_CHARS:
        raise ValueError("文档解析后的文字总量超过处理上限")
    return chunks
